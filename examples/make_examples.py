"""Generate a content-free example template and spreadsheet.

The repository deliberately ships no real documents. Running this script writes
a small, entirely invented pair of files into ``examples/generated`` so the
program and its tests can be exercised end to end:

    python examples/make_examples.py

The same function is used by the test suite, which is why the output is plain
and predictable rather than decorative.
"""
from __future__ import annotations

from pathlib import Path

OUTPUT = Path(__file__).resolve().parent / "generated"

# Caption -> value. Captions are what the matcher reads; values are invented.
SAMPLE_ROWS = [
    ("Customer Name", "Northwind Trading Ltd."),
    ("Customer Address", "19 Harbour Road, Bristol"),
    ("Contact Person", "Alex Morgan"),
    ("Order Number", "ORD-2026-0184"),
    ("Issue Date", "04.03.2026"),
    ("Product", "Industrial Drying Unit"),
    ("Quantity", 12),
    ("Unit Price", 1450.5),
    ("Total Amount", 17406.0),
    ("Delivery Term", "Ex Works"),
]

TEMPLATE_BODY = [
    ("heading", "Order Confirmation"),
    ("paragraph", "Order number: {{ order_number }}"),
    ("paragraph", "Issue date: {{ issue_date }}"),
    ("paragraph", ""),
    ("paragraph", "Customer: {{ customer_name }}"),
    ("paragraph", "Address: {{ customer_address }}"),
    ("paragraph", "Contact person: {{ contact_person }}"),
    ("paragraph", ""),
    ("paragraph", "We confirm the following order."),
]

TEMPLATE_TABLE = [
    ("Product", "{{ product }}"),
    ("Quantity", "{{ quantity }}"),
    ("Unit price", "{{ unit_price }}"),
    ("Total amount", "{{ total_amount }}"),
    ("Delivery term", "{{ delivery_term }}"),
    ("Reference", "{{ internal_reference }}"),
]


def make_template(path: Path) -> Path:
    """Write a small Word template containing placeholders."""
    from docx import Document

    document = Document()
    for kind, text in TEMPLATE_BODY:
        if kind == "heading":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)

    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for caption, placeholder in TEMPLATE_TABLE:
        cells = table.add_row().cells
        cells[0].text = caption
        cells[1].text = placeholder

    document.add_paragraph("")
    document.add_paragraph("Issued by {{ issued_by }}.")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def make_workbook(path: Path) -> Path:
    """Write a spreadsheet laid out as caption/value pairs."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Order"
    sheet["A1"] = "Order Details"
    for index, (caption, value) in enumerate(SAMPLE_ROWS, start=3):
        sheet.cell(row=index, column=1, value=caption)
        sheet.cell(row=index, column=2, value=value)

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(path))
    return path


def build(folder: Path | None = None) -> tuple[Path, Path]:
    """Create both example files and return their paths."""
    folder = Path(folder) if folder else OUTPUT
    template = make_template(folder / "order_confirmation.docx")
    data = make_workbook(folder / "order_data.xlsx")
    return template, data


if __name__ == "__main__":
    template_path, data_path = build()
    print(f"Template : {template_path}")
    print(f"Data     : {data_path}")
    print("\nIn the program choose the folder above as the template folder,")
    print("then choose the spreadsheet as the data source.")
