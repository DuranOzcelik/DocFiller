"""Tests for everything that does not need a display or Office.

The data folder is redirected to a temporary directory, so running the tests
never touches the real accounts, history or settings.
"""
from __future__ import annotations

import ast
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = Path(tempfile.mkdtemp(prefix="docfiller-data-"))
os.environ["APPDATA"] = str(DATA)
os.environ["LOCALAPPDATA"] = str(DATA)
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from docfiller import (  # noqa: E402
    accounts,
    document_sets,
    drafts,
    fields,
    formatting,
    history,
    i18n,
    library,
    memory,
    renderer,
    settings,
    sheet_fill,
    templates,
    workbook,
)
from docfiller.engine import NoDataSource, resolve  # noqa: E402
from examples.make_examples import build  # noqa: E402


def setUpModule():
    account = accounts.create("Test User", role="Tester")
    accounts.sign_in(account)


def tearDownModule():
    shutil.rmtree(DATA, ignore_errors=True)


class ExampleFiles(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = Path(tempfile.mkdtemp(prefix="docfiller-test-"))
        cls.template_path, cls.data_path = build(cls.folder)
        cls.template = templates.read(cls.template_path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.folder, ignore_errors=True)

    def resolved(self):
        return resolve([self.data_path], [self.template])


class TestTemplates(ExampleFiles):
    def test_placeholders_are_found(self):
        self.assertTrue(self.template.ok, self.template.error)
        for name in ("customer_name", "order_number", "total_amount"):
            self.assertIn(name, self.template.placeholders)

    def test_discover(self):
        self.assertIn(self.template.name, [t.name for t in templates.discover(self.folder)])


class TestReading(ExampleFiles):
    def test_label_values_are_read(self):
        captions = {lv.label for lv in workbook.read_source(self.data_path)}
        self.assertIn("Customer Name", captions)

    def test_resolve_matches_and_keeps_values_whole(self):
        project = self.resolved().project
        self.assertEqual(project.text("customer_name"), "Northwind Trading Ltd.")
        self.assertEqual(project.origin("customer_name"), "source")

    def test_no_data_source(self):
        with self.assertRaises(NoDataSource):
            resolve([], [self.template])

    def test_typed_value_wins(self):
        project = self.resolved().project
        project.type_value("customer_name", "Typed Name")
        self.assertEqual(project.text("customer_name"), "Typed Name")
        project.clear_typed("customer_name")
        self.assertEqual(project.text("customer_name"), "Northwind Trading Ltd.")


class TestFilling(ExampleFiles):
    def test_fill_writes_values(self):
        from docx import Document

        project = self.resolved().project
        out = Path(tempfile.mkdtemp(dir=self.folder))
        target = renderer.fill(self.template, project, out)
        text = "\n".join(p.text for p in Document(str(target)).paragraphs)
        self.assertIn("Northwind Trading Ltd.", text)
        self.assertNotIn("{{", text)

    def test_missing_fields_are_listed(self):
        project = self.resolved().project
        names = [f.name for f in fields.missing(self.template, project)]
        self.assertIn("internal_reference", names)


class TestDrafts(ExampleFiles):
    def test_only_the_user_difference_is_stored(self):
        fresh = self.resolved().project
        current = self.resolved().project
        current.type_value("internal_reference", "REF-1")
        record = drafts.capture(fresh, current)
        self.assertEqual(list(record["typed"]), ["internal_reference"])
        key = drafts.identity([self.data_path])
        drafts.save(key, record)
        again = self.resolved().project
        drafts.apply(again, drafts.load(key))
        self.assertEqual(again.text("internal_reference"), "REF-1")


class TestSheetFill(ExampleFiles):
    def test_fields_are_found(self):
        found, _ = sheet_fill.extract_fields(self.data_path)
        self.assertTrue(found)


class TestFormatting(unittest.TestCase):
    def test_both_number_conventions(self):
        self.assertEqual(formatting.to_number("1.234,5"), 1234.5)
        self.assertEqual(formatting.to_number("1,234.5"), 1234.5)

    def test_single_line(self):
        self.assertEqual(formatting.single_line("a\n b"), "a b")


class TestStores(unittest.TestCase):
    def test_accounts(self):
        account = accounts.create("Second Person", "secret")
        self.assertTrue(accounts.password_matches(account, "secret"))
        self.assertFalse(accounts.password_matches(account, "wrong"))
        self.assertTrue(accounts.remove(account.id))

    def test_library(self):
        library.save(library.Record("Acme", {"customer_name": "Acme Ltd."}))
        self.assertEqual(library.find("acme").values["customer_name"], "Acme Ltd.")
        self.assertTrue(library.delete("Acme"))

    def test_document_sets(self):
        saved = document_sets.save(document_sets.DocumentSet(id="", name="Order pack", templates=["a.docx"]))
        self.assertEqual(document_sets.required_templates(saved.id), ["a.docx"])
        self.assertTrue(document_sets.delete(saved.id))

    def test_memory(self):
        memory.remember("city", "Bristol")
        self.assertIn("Bristol", memory.values("city"))
        memory.clear()
        self.assertEqual(memory.values("city"), [])

    def test_history_empty_and_clear(self):
        history.clear()
        self.assertEqual(history.entries(), [])

    def test_settings(self):
        settings.set_machine(mode="dark")
        self.assertEqual(settings.mode(), "dark")


class TestTranslations(unittest.TestCase):
    def test_every_interface_string_has_turkish(self):
        keys = set()
        for file in (ROOT / "src" / "docfiller").rglob("*.py"):
            for node in ast.walk(ast.parse(file.read_text(encoding="utf-8"))):
                if (isinstance(node, ast.Call) and getattr(node.func, "id", None) == "tr"
                        and node.args and isinstance(node.args[0], ast.Constant)):
                    keys.add(node.args[0].value)
        self.assertEqual(sorted(k for k in keys if k not in i18n.TURKISH), [])


if __name__ == "__main__":
    unittest.main()
