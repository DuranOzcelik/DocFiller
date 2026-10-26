"""Checking a produced document before it counts as produced.

Two things are checked:

1. Is any unfilled ``{{ placeholder }}`` or ``{% ... %}`` left in the document?
   That means a variable never reached the template, and it must not appear in
   a delivered document.
2. Is a field the document set marks as required still empty?

A document with a problem is still produced — so it can be fixed by hand — but
it is not presented as a success: the warning names the exact field.

The text is not read through ``paragraph.text``: placeholders inside content
controls and text boxes are only visible to a deep search over every ``w:t``.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.oxml.ns import qn

from . import fields, templates

LEFTOVER = re.compile(r"\{\{.*?\}\}|\{%.*?%\}", re.DOTALL)


@dataclass
class Check:
    document: str                     # template name
    file: Path | None = None
    leftovers: list[str] = field(default_factory=list)
    empty_required: list[str] = field(default_factory=list)
    error: str | None = None

    @property
    def passed(self) -> bool:
        return not (self.leftovers or self.empty_required or self.error)

    def message(self) -> str:
        if self.error:
            return f"could not be checked: {self.error}"
        parts = []
        if self.leftovers:
            shown = ", ".join(self.leftovers[:4])
            if len(self.leftovers) > 4:
                shown += f" (+{len(self.leftovers) - 4})"
            parts.append(f"unfilled placeholder left: {shown}")
        if self.empty_required:
            parts.append("required field empty: " + ", ".join(templates.label(n) for n in self.empty_required))
        return "; ".join(parts) if parts else "no problems found"


def _deep_text(document) -> str:
    return "".join(t.text or "" for t in document.element.body.findall(".//" + qn("w:t")))


def leftovers(document) -> list[str]:
    """Unfilled placeholders left in an OPEN document (deduplicated, in order).

    A placeholder may be split across runs, so the search runs over the joined
    text of the whole document rather than run by run.
    """
    return list(dict.fromkeys(m.group(0).strip() for m in LEFTOVER.finditer(_deep_text(document))))


def check(template: templates.Template, file: Path, project, document=None) -> Check:
    """Validate one produced document. `document`, when given, is the object
    still open from filling, so the file is not read back from disk."""
    result = Check(document=template.name, file=Path(file))
    try:
        opened = document if document is not None else docx.Document(str(file))
        result.leftovers = leftovers(opened)
    except Exception as exc:  # noqa: BLE001
        result.error = str(exc)
        return result
    result.empty_required = [f.name for f in fields.missing(template, project) if f.required]
    return result
